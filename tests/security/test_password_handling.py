"""Regression coverage for encryption classification and HTTP rejection."""
import asyncio
import hashlib
import json

import httpx
import msoffcrypto
import pytest
import xlrd
from docx import Document
from openpyxl import Workbook

from ingestion_pipeline import api
from ingestion_pipeline.extraction import excel
from ingestion_pipeline.processing.preflight import PreflightError, check_document


@pytest.mark.parametrize("suffix", [".xlsx", ".docx"])
def test_real_encrypted_ooxml_requires_password(tmp_path, suffix):
    plain = tmp_path / ("plain" + suffix)
    if suffix == ".xlsx":
        book = Workbook()
        book.active.append(["single row"])
        book.save(plain)
        book.close()
    else:
        doc = Document()
        doc.add_paragraph("content")
        doc.save(plain)
    check_document(plain)
    encrypted = tmp_path / ("encrypted" + suffix)
    with plain.open("rb") as source, encrypted.open("wb") as target:
        msoffcrypto.OfficeFile(source).encrypt("test-password", target)
    with pytest.raises(PreflightError) as error:
        check_document(encrypted)
    assert error.value.code == "password_required"


@pytest.mark.parametrize("suffix", [".xlsx", ".docx"])
@pytest.mark.parametrize("tail", [b"", b"0" * 64])
def test_truncated_ole_is_corruption(tmp_path, suffix, tail):
    path = tmp_path / ("broken" + suffix)
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + tail)
    with pytest.raises(PreflightError) as error:
        check_document(path)
    assert error.value.code == "corrupt_document"


@pytest.mark.parametrize("inspection_fails", [False, True])
@pytest.mark.parametrize("message,expected", [
    ("Workbook is encrypted", "password_required"),
    ("Invalid workbook", "corrupt_document"),
    ("Invalid password metadata", "corrupt_document"),
])
def test_xls_reader_error_classification(tmp_path, monkeypatch, inspection_fails, message, expected):
    path = tmp_path / "workbook.xls"
    path.write_bytes(b"fixture")

    def fail_inspection(*args, **kwargs):
        raise ValueError("Invalid Office metadata")

    if inspection_fails:
        monkeypatch.setattr(msoffcrypto, "OfficeFile", fail_inspection)
    else:
        monkeypatch.setattr(excel, "msoffcrypto", None)

    def fail_read(*args, **kwargs):
        raise xlrd.biffh.XLRDError(message)

    monkeypatch.setattr(xlrd, "open_workbook", fail_read)
    with pytest.raises(PreflightError) as error:
        excel.parse_excel(str(path))
    assert error.value.code == expected


@pytest.mark.parametrize("route", ["/extract", "/validate-attachments"])
@pytest.mark.parametrize("code", ["password_required", "corrupt_document"])
def test_sync_routes_return_structured_rejections(monkeypatch, route, code):
    token = "test-password-handling-api-token".ljust(40, "x")
    monkeypatch.setenv("INGESTION_API_KEY_SHA256_JSON", json.dumps({
        hashlib.sha256(token.encode()).hexdigest(): {
            "tenant_id": "test", "subject": "owner", "roles": ["extract"],
        },
    }))
    monkeypatch.setattr(api, "feature_enabled", lambda name: True)

    def reject(**kwargs):
        raise PreflightError(code, "Document rejected")

    monkeypatch.setattr(api, "process_email_attachments", reject)

    async def request():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app), base_url="http://test") as client:
            return await client.post(route, headers={"Authorization": "Bearer " + token},
                                     files={"files": ("input.xls", bytes.fromhex("d0cf11e0a1b11ae1"), "application/vnd.ms-excel")})

    response = asyncio.run(request())
    assert response.status_code == 400
    assert response.json() == {"error": code, "detail": "Document rejected"}

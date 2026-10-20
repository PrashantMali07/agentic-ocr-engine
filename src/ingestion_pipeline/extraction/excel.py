"""Bounded XLSX/XLS extraction into the unified table-block schema."""

import logging
from contextlib import closing
from html import escape
from pathlib import Path

from ..processing.preflight import PreflightError

log = logging.getLogger("ingest.excel")

MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 10_000
MAX_COLUMNS = 100
MAX_TOTAL_CELLS = 100_000
MAX_CELL_CHARS = 10_000
MAX_HTML_BYTES = 4 * 1024 * 1024

try:
    import msoffcrypto
except ImportError:  # pragma: no cover - dependency is in the supported images
    msoffcrypto = None


def _is_encrypted(path: Path) -> bool:
    if msoffcrypto is None:
        return False
    try:
        with path.open("rb") as source:
            return msoffcrypto.OfficeFile(source).is_encrypted()
    except Exception:
        return False


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        raise PreflightError("cell_limit", "Spreadsheet cell exceeds the supported limit")
    return escape(text, quote=True)


def _render_sheet(rows, name: str, budget: list[int]) -> str | None:
    html = ["<table><caption>", escape(name, quote=True), "</caption>"]
    has_data = False
    for row_number, row in enumerate(rows, start=1):
        if row_number > MAX_ROWS_PER_SHEET:
            raise PreflightError("sheet_limit", "Spreadsheet has too many rows")
        cells = list(row)
        if len(cells) > MAX_COLUMNS:
            raise PreflightError("sheet_limit", "Spreadsheet has too many columns")
        budget[0] += len(cells)
        if budget[0] > MAX_TOTAL_CELLS:
            raise PreflightError("sheet_limit", "Spreadsheet has too many cells")
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        has_data = True
        line = "<tr>" + "".join(f"<td>{_cell_text(value)}</td>" for value in cells) + "</tr>"
        budget[1] += len(line.encode("utf-8"))
        if budget[1] > MAX_HTML_BYTES:
            raise PreflightError("output_limit", "Spreadsheet output exceeds the supported limit")
        html.append(line)
    html.append("</table>")
    return "".join(html) if has_data else None


def _xlsx_sheets(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise PreflightError("sheet_limit", "Spreadsheet has too many sheets")
        for sheet in workbook.worksheets:
            if (sheet.max_row and sheet.max_row > MAX_ROWS_PER_SHEET) or (sheet.max_column and sheet.max_column > MAX_COLUMNS):
                raise PreflightError("sheet_limit", "Spreadsheet dimensions exceed the supported limit")
            yield sheet.title, sheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def _xls_sheets(path: Path):
    import xlrd

    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        # xlrd can recognize FILEPASS even when the Office pre-check cannot.
        if str(exc) == "Workbook is encrypted":
            raise PreflightError("password_required", "Document requires a password") from exc
        raise
    try:
        if workbook.nsheets > MAX_SHEETS:
            raise PreflightError("sheet_limit", "Spreadsheet has too many sheets")
        for index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(index)
            if sheet.nrows > MAX_ROWS_PER_SHEET or sheet.ncols > MAX_COLUMNS:
                raise PreflightError("sheet_limit", "Spreadsheet dimensions exceed the supported limit")
            yield sheet.name, (sheet.row_values(row) for row in range(sheet.nrows))
    finally:
        workbook.release_resources()


def parse_excel(file_path: str, image_out_dir: str = "./extracted_images") -> dict:
    path = Path(file_path)
    if not path.is_file():
        raise PreflightError("missing_document", "Spreadsheet input is missing")
    if _is_encrypted(path):
        raise PreflightError("password_required", "Document requires a password")

    record = {
        "source_file": str(path), "format": "excel", "blocks": [], "extracted_images": [],
    }
    budget = [0, 0]  # cells, rendered HTML bytes across all sheets
    try:
        sheets = _xls_sheets(path) if path.suffix.lower() == ".xls" else _xlsx_sheets(path)
        with closing(sheets):
            for name, rows in sheets:
                html = _render_sheet(rows, name, budget)
                if html is None:
                    continue
                record["blocks"].append({
                    "block_label": "table", "content": html, "order": len(record["blocks"]),
                    "page": None, "confidence": 1.0, "source_file": path.name, "format": "excel",
                })
    except PreflightError:
        raise
    except Exception as exc:
        log.warning("Spreadsheet parser rejected %s", path.name)
        raise PreflightError("corrupt_document", "Spreadsheet cannot be parsed") from exc
    return record
